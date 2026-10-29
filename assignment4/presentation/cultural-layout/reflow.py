from pathlib import Path
from copy import deepcopy
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
BASE=Path('assignment4/presentation/selected-styles.pptx')
OUT=Path('assignment4/presentation/cultural-layout/ICT-layout.pptx')
p=Presentation(BASE); original=Presentation(BASE)
INK='282043'; PURPLE='7044DA'; WHITE='FFFFFF'; LAV='EEE7FC'; PINK='FFDEED'; MINT='D6F4E5'
changed=[]; frames=[]
def rect(s,x,y,w,h,fill=LAV):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.adjustments[0]=.06
 style=sh._element.find(qn('p:style'))
 if style is not None:sh._element.remove(style)
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background();return sh

def text(s,t,x,y,w,h,size=26,bold=False,color=INK,align=PP_ALIGN.LEFT):
 h+=.18
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
 for i,line in enumerate(t.split('\n')):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.alignment=align;para.line_spacing=1.08;para.space_before=Pt(0);para.space_after=Pt(6)
  r=para.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 frames.append((len(changed) and changed[-1],t,x,y,w,h))
 return sh

def start(n,title=None,right=False):
 s=p.slides[n-1];changed.append(n)
 for sh in list(s.shapes)[1:]:sh._element.getparent().remove(sh._element)
 if title:text(s,title,.95,.5,18.1,.75,35,True,WHITE,PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)
 return s

def block(s,label,body,x,y,w,h,fill=LAV,sz=26):
 rect(s,x,y,w,h,fill);text(s,label,x+.35,y+.3,w-.7,.85,29,True,PURPLE);text(s,body,x+.35,y+1.1,w-.7,h-1.4,sz)

def note(s,t,x,y,w,h,size=20):
 rect(s,x,y,w,h,WHITE);text(s,t,x+.25,y+.2,w-.5,h-.35,size)

# Keep the approved separated-cell table, but place it to the right of a clear introduction.
s=start(2)
text(s,'Part A',1,1.8,5,.7,32,True,WHITE)
text(s,'Multimedia\nformats',1,2.65,5.3,1.9,48,True,WHITE)
block(s,'In our prototype','WebP for the attendance preview; inline SVG for interface icons.',1,5.05,5.35,2.8,WHITE,24)
data=[['FORMAT','KEY TRAIT','BEST FIT'],['PNG','Lossless; supports transparency','Screenshots and graphics where crisp text or edges matter'],['JPEG','Lossy; compact for photographs','Photos and images where a smaller file is more important than perfect detail'],['WebP','Lossy or lossless; compact','Web images and the attendance preview used in our prototype']]
x0=6.9;y=1.8;widths=[1.6,4.0,5.35]
for i,row in enumerate(data):
 x=x0;h=.65 if i==0 else 1.7
 for j,t in enumerate(row):
  rect(s,x,y,widths[j],h,PURPLE if i==0 else [LAV,PINK,MINT][i-1] if j==0 else WHITE)
  text(s,t,x+.16,y+.14,widths[j]-.32,h-.24,18 if i==0 else 23,i==0 or j==0,WHITE if i==0 else INK)
  x+=widths[j]+.18
 y+=h+.2

# Reference page 5: two offset numeric callouts, with supporting text beneath.
s=start(4,'A2. Measured experiment')
rect(s,1,1.8,8.6,2.45,WHITE);text(s,'PNG · original',1.35,2.08,7.9,.45,23,True,PURPLE);text(s,'58.39 KiB',1.35,2.65,7.9,.95,55,True)
text(s,'attendance-preview.png · 390 × 844 px',1.35,3.65,7.9,.4,20)
rect(s,10.35,2.4,8.6,2.45,MINT);text(s,'WebP Q80 · in app',10.7,2.68,7.9,.45,23,True,PURPLE);text(s,'15.75 KiB',10.7,3.25,7.9,.95,55,True)
text(s,'attendance-preview.webp · 390 × 844 px',10.7,4.25,7.9,.4,20)
rect(s,1,4.53,8.6,.82,PURPLE);text(s,'42.64 KiB saved · 73.0% smaller',1.3,4.72,8,.5,26,True,WHITE)
block(s,'JPEG exports · pending','attendance-preview.jpg — JPEG Q80\nattendance-preview-low.jpg — JPEG Q40\nBoth: 390 × 844 px. Files not supplied; export and inspect.',1,5.8,8.6,2.7,WHITE,21)
block(s,'10-second clip · pending','Version 1: proposed MP4 H.264, 720p.\nVersion 2: same clip, lower video bitrate.\nFiles not supplied; compare text and motion.',10.35,5.25,8.6,3.25,WHITE,22)
note(s,'Visual quality checks are still pending. JPEG and clip entries are a test plan, not completed measurements.\nSource: supplied project inventory, 23 Oct 2026. 1 KiB = 1024 bytes.',1,8.7,17.95,.94,17)

# Reference page 3: dominant explanatory block left and a smaller supporting block right.
s=start(5,'A2. Visual quality check')
block(s,'Compare like for like','Use the same attendance screen in all four exports. Compare an identical crop at 100% zoom so text edges and calendar lines remain comparable.',1,2,10.4,4.3,WHITE,31)
block(s,'Four exports','PNG\nJPEG Q80\nJPEG Q40\nWebP Q80',12,2.8,6.95,4.25,LAV,30)
block(s,'Evidence still needed','One side-by-side screenshot, JPEG file sizes, and two exports of the same 10-second clip. Rate each result as sharp, acceptable or poor after inspection.',1,6.65,10.4,2.55,WHITE,22)
note(s,'Comparison images and recording are not included in this deck. Quality ratings remain pending.',12,7.4,6.95,1.75,22)

# Reference page 8: substantial text on the left, headline / main number on the right.
s=start(6,'A2. Our format choices',right=True)
block(s,'For the My AITU / LMS bridge','WebP Q80 supplies the attendance preview at 390 × 844 px. PNG remains the original evidence image because it preserves text and edges without lossy compression.\nInline SVG keeps icons sharp at different sizes and shares the HTML response. CSS provides the login pulse without a separate video download.',1,1.8,10.1,4.65,WHITE,25)
block(s,'Planned demonstration','MP4 with H.264 video and AAC audio for broad playback support. JPEG settings and the ten-second clip still need measured comparison before we can report their quality or size.',1,6.8,10.1,2.6,WHITE,23)
text(s,'73%',12,2.15,6.9,1.6,96,True,WHITE)
text(s,'smaller preview file',12.05,3.95,6.8,.7,31,True,WHITE)
block(s,'Same pixels. Less data.','PNG 58.39 KiB → WebP 15.75 KiB\n42.64 KiB saved at the same resolution.',11.75,5.1,7.2,2.65,MINT,25)
note(s,'Measured file sizes only.\nNetwork speed needs its own test.',11.75,8.1,7.2,1.3,22)

# Reference page 4: larger narrative left, two smaller staggered modules right.
s=start(9,'B2. Integration and optimisation')
block(s,'What is integrated','WebP preview: 390 × 844 px, displayed at 250 px wide.\n18 inline SVG icons and a CSS login pulse.\nReduced-motion preference disables animation.\nThe full page still needs a 360 px browser check.',1,1.85,8.3,4.9,WHITE,25)
rect(s,9.95,2.35,9,2.15,MINT);text(s,'−73.0%',10.3,2.7,3.3,.85,43,True,PURPLE);text(s,'Preview image',14,2.65,4.5,.5,24,True);text(s,'59,787 B → 16,126 B',14,3.35,4.5,.7,25)
rect(s,10.7,4.85,8.25,2.15,LAV);text(s,'−30.8%',11.05,5.2,3.3,.85,43,True,PURPLE);text(s,'4 front-end files + preview',14.6,5.1,3.9,.8,22,True);text(s,'141,601 B → 97,941 B',14.6,6.03,3.9,.65,22)
block(s,'DevTools measurements: still required','Disable cache, select Fast 4G and reload the same screen before and after. Capture transferred bytes, request count and Finish/Load time. Image-request delta should be 0; verify the total.',1,7.15,11.7,2.4,WHITE,21)
note(s,'Controlled local payload excludes headers, compression, API, fonts and map traffic. Stage 1: <3 s per request at 1,000 users remains unverified.',13.05,7.4,5.9,2.15,20)

# Reference page 7: offset paired panels, followed by a slim contextual note.
s=start(10,'B3. Accessibility of media',right=True)
block(s,'Meaningful images have alt text','The test attendance preview describes the month and status counts. Decorative SVGs use aria-hidden="true"; controls retain text or accessible labels.',1,1.8,9,3.0,WHITE,24)
rect(s,1,5.1,9,2.75,LAV)
text(s,'<img src="/attendance-preview.webp"\nwidth="390" height="844"\nalt="Тестовый экран посещаемости: календарь за октябрь 2026 года и счётчики — 0 присутствий, 1 пропуск, 0 опозданий, 1 запись без отметки.">',1.3,5.4,8.4,2.15,20)
block(s,'Motion follows user preference','Progress is visible as text and a pulse. Animation and transitions are disabled when reduced motion is requested.',10.6,2.65,8.35,2.7,WHITE,24)
rect(s,10.6,5.65,8.35,2.2,LAV)
text(s,'@media (prefers-reduced-motion: reduce) {\n  *, *::before, *::after { animation: none !important;\n  transition: none !important;\n  scroll-behavior: auto !important; } }',10.9,5.95,7.75,1.6,19)
note(s,'No speech media or sound playback: captions and sound equivalents are currently N/A. No live text sits over raster images, so the 4.5:1 text-over-image test is N/A. Source: supplied part_b.md, B3.',1,8.3,17.95,1.15,20)

# Reference page 2: a stepped sequence, with alternation rather than tabular columns.
s=start(12,'C1. Storyboard')
steps=[('01  ·  0:00–0:30','Login screen','Nurassyl: how the login process looks',1,1.8),('02  ·  0:30–1:10','Main pages, burger menu and references from Stage 1','Salima: website content and design decisions',9.35,3.05),('03  ·  1:10–1:40','Statistics and survey from Stage 1','Azhar: the purpose of the project',1,5.5),('04  ·  1:40–2:10','Demo sketches','Azhar: the future of the project',9.35,6.75)]
for i,(label,screen,narr,x,y) in enumerate(steps):
 rect(s,x,y,8,2.45,[WHITE,LAV,MINT,PINK][i]);text(s,label,x+.3,y+.22,7.4,.5,25,True,PURPLE);text(s,screen,x+.3,y+.88,7.4,.8,24,True);text(s,narr,x+.3,y+1.85,7.4,.48,20)

# Keep the native outlined bubble and landscape, but give its instructions a readable hierarchy.
s=start(16,'D1. Updated prototype')
text(s,'Run the prototype',3.9,1.9,11.8,.55,31,True,WHITE)
rect(s,3.8,2.58,12,.65,WHITE)
sh=text(s,'github.com/d30mantcoder/ICTProject',4,2.7,11.6,.48,25,False,PURPLE)
sh.text_frame.paragraphs[0].runs[0].hyperlink.address='https://github.com/d30mantcoder/ICTProject'
text(s,'01  Clone the repository\n02  Create a virtual environment',3.9,3.45,5.7,1.4,25,False,WHITE)
text(s,'03  Install requirements.txt\n04  Run bridge-test/app.py with Python',10,3.45,5.9,1.4,25,False,WHITE)

# Reference page 4: one large panel left, two concise sections right; no changelog grid.
s=start(17,'D2. Change log',right=True)
block(s,'01  Real LMS data','Added fetching from LMS to make the website usable.\nStatus: using real LMS data.',1,2,8.5,4.6,WHITE,31)
block(s,'02  A more usable interface','The original LMS UX was uncomfortable. The redesigned interface places the most useful buttons on the main page.',10.2,2.75,8.75,2.75,LAV,25)
block(s,'03  Multi-user sessions','Previously, more than one person could not work at the same time. Asynchronous sessions now let the server handle multiple users.',10.2,5.85,8.75,3.05,MINT,25)
text(s,'Data → interface → concurrent use',1.1,7.1,8.3,1.2,31,True,WHITE)

for n in changed:
 assert E.tostring(original.slides[n-1].shapes[0]._element)==E.tostring(p.slides[n-1].shapes[0]._element),n
p.save(OUT)
import json
OUT.with_suffix('.frames.json').write_text(json.dumps(frames,ensure_ascii=False,indent=2))
print('Saved',OUT,'changed',changed,'all original backgrounds preserved')
