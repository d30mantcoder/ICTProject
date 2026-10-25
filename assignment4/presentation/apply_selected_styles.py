from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from collections import Counter
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,MSO_AUTO_SIZE,PP_ALIGN
from pptx.oxml.ns import qn
BASE=Path('assignment4/presentation/restored-text-fixed.pptx')
OUT=Path('assignment4/presentation/selected-styles.pptx')
p=Presentation(BASE);changed={2,3,4,7,8,9,12,17,18}
WHITE='FFFFFF';INK='282043';PURPLE='7044DA';PASTELS=['E3D4FF','FFDEED','D6F4E5']
def remove(s,sh):sh._element.getparent().remove(sh._element)
def cell(s,text,x,y,w,h,size=23,fill=WHITE,bold=False,color=INK,center=False,pad=.15,radius=.14):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.adjustments[0]=min(.18,radius/min(w,h))
 style=sh._element.find(qn('p:style'))
 if style is not None:sh._element.remove(style)
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 tf=sh.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 tf.margin_left=tf.margin_right=Inches(pad);tf.margin_top=tf.margin_bottom=Inches(.015 if h<.6 else .08)
 para=tf.paragraphs[0];para.alignment=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT;para.space_before=para.space_after=Pt(0);para.line_spacing=1.05
 r=para.add_run();r.text=text;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 return sh

def grid(s,data,x,y,widths,rh,hh=.65,gap=.1,size=22,header_size=None,boldcol=0):
 assert all(len(row)==len(widths) for row in data)
 for i,row in enumerate(data):
  h=hh if i==0 else rh[i-1] if isinstance(rh,list) else rh
  cx=x
  for j,t in enumerate(row):
   cell(s,t,cx,y,widths[j],h,size=(header_size or size) if i==0 else size,fill=PURPLE if i==0 else PASTELS[(i-1)%3] if j==0 else WHITE,bold=i==0 or j==boldcol,color=WHITE if i==0 else INK,center=(widths[j]<.8),pad=.08 if size<20 else .19,radius=.07 if h<.7 else .15)
   cx+=widths[j]+gap
  y+=h+gap
 return y-gap

def flatten_text(sh):
 if sh.shape_type==6:return ' | '.join(filter(None,(flatten_text(c) for c in sh.shapes)))
 return sh.text if sh.has_text_frame else ''

# Slide 2 — the selected 09 treatment, using its existing exact cell content.
s=p.slides[1];old=next(sh for sh in s.shapes if sh.has_table);data=[[c.text for c in row.cells] for row in old.table.rows];remove(s,old)
grid(s,data,.95,1.75,[2.9,6.1,8.6],1.29,hh=.78,gap=.25,size=25,header_size=20)

# Previously imported tables consist of individual text boxes. Read all cells before replacing their visual structure.
configs={
 3:dict(start=3,end=96,cols=7,widths=[2.1,1.15,2.45,2.65,2.35,1.8,5.0],rh=.425,hh=.6,gap=.055,size=17.25,footer=7.32,boldcol=0),
 4:dict(start=3,end=40,cols=4,widths=[4.8,5.2,2.25,5.65],rh=.62,hh=.6,gap=.1,size=20,footer=6.78,boldcol=0),
 7:dict(start=3,end=84,cols=5,widths=[.63,3.65,5.5,4.0,4.15],rh=.4,hh=.52,gap=.055,size=17.25,footer=7.82,boldcol=1),
 8:dict(start=3,end=105,cols=6,widths=[.63,3.4,3.15,4.4,2.05,4.245],rh=.37,hh=.5,gap=.045,size=16.5,footer=7.77,boldcol=1),
 9:dict(start=8,end=25,cols=4,widths=[7.15,3.55,3.55,3.61],rh=.48,hh=.52,gap=.08,size=20,footer=None,boldcol=0),
}
for n,c in configs.items():
 s=p.slides[n-1];shapes=list(s.shapes);texts=[sh.text for sh in shapes[c['start']:c['end']+1] if sh.has_text_frame and sh.text.strip()]
 assert len(texts)%c['cols']==0,(n,len(texts))
 data=[texts[i:i+c['cols']] for i in range(0,len(texts),c['cols'])]
 for sh in shapes[c['start']:c['end']+1]:remove(s,sh)
 if c['footer'] is not None:
  panel=shapes[2];bottom=(panel.top+panel.height)/914400;panel.top=Inches(c['footer']);panel.height=Inches(bottom-c['footer'])
 grid(s,data,.95,4.4 if n==9 else 1.75,c['widths'],c['rh'],hh=c['hh'],gap=c['gap'],size=c['size'],header_size=c['size'] if n!=8 else 16,boldcol=c['boldcol'])

# Storyboard: retain the original narrow composition and all text.
s=p.slides[11];shapes=list(s.shapes);data=[['Duration','What is on screen','Narration'],['0:00-0:30','Login screen','Nurassyl: How login proccess looks'],['0:30-1:10','Main pages, burger-menu, references from stage 1','Salima: Website content and design desecions'],['1:10-1:40','Statistics, survey from stage 1','Azhar: the purpose of project'],['1:40-2:10','Demo sketches','Azhar: Future of the project']]
for sh in shapes[2:]:remove(s,sh)
grid(s,data,.89,1.75,[2.2,4.8,4.7],1.6,hh=.8,gap=.15,size=23,header_size=22)

# Changelog.
s=p.slides[16];shapes=list(s.shapes);data=[['Change','Reason','Status'],['Added fetching from LMS','To make the website usable','Using real LMS data'],['Redesign of UX of LMS','UX of original LMS was uncomfortable','Most useful buttons takes place on the main page'],['Added asynchronius sessions','It was impossible to work for more than 1 person at the same time before','Server handles multi-user sessions']]
for sh in shapes[2:]:remove(s,sh)
grid(s,data,.95,2.25,[5.5,6.0,6.3],[1.4,1.8,1.8],hh=.8,gap=.15,size=25)

# Team: the approved variant 17, without option labels.
s=p.slides[17];old=next(sh for sh in s.shapes if sh.has_table);team=[[c.text for c in row.cells] for row in list(old.table.rows)[1:]];remove(s,old)
head=s.shapes[1];head.top=Inches(1.38);head.text_frame.margin_top=Inches(.04)
for i,(name,parts,contribution) in enumerate(team):
 x=.95+i*6.15
 card=cell(s,'',x,2.73,5.8,4.85,fill=WHITE,radius=.23)
 badge=cell(s,parts,x+.4,3.12,1.45,.5,19,fill=[ 'EEE7FC','FFE1EE','D8F5E8'][i],bold=True,color=PURPLE,center=True,pad=.08,radius=.22)
 def label(t,xx,yy,ww,hh,size,color=INK,bold=False):
  sh=s.shapes.add_textbox(Inches(xx),Inches(yy),Inches(ww),Inches(hh));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
  para=tf.paragraphs[0];para.space_after=Pt(0);r=para.add_run();r.text=t;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 label(name,x+.4,4.01,5,.65,34,bold=True)
 rule=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.4),Inches(5.05),Inches(5),Inches(.012));rule.fill.solid();rule.fill.fore_color.rgb=RGBColor.from_string('DED5F0');rule.line.fill.background()
 style=rule._element.find(qn('p:style'))
 if style is not None:rule._element.remove(style)
 label('CONTRIBUTION',x+.4,5.4,5,.3,17,'716785',True);label(contribution,x+.4,6.03,5,.95,27)

# Compare wording, retaining every existing phrase; only repeated card labels are new.
def texts(s):
 result=[]
 def scan(sh):
  if sh.shape_type==6:
   for c in sh.shapes:scan(c)
  elif sh.has_table:
   for row in sh.table.rows:
    for c in row.cells:
     if c.text.strip():result.append(c.text.strip())
  elif sh.has_text_frame and sh.text.strip():result.append(sh.text.strip())
 for sh in s.shapes:scan(sh)
 return Counter(result)
original=Presentation(BASE)
for n in changed:
 before=texts(original.slides[n-1]);after=texts(p.slides[n-1]);missing=before-after
 if n==18:
  for key in ['Members','Parts','Contribution']:missing.pop(key,None)
 assert not missing,(n,missing)
 assert E.tostring(original.slides[n-1].shapes[0]._element)==E.tostring(p.slides[n-1].shapes[0]._element),n
p.save(OUT)
# Keep untouched slide XML byte-for-byte identical to the current original.
temp=OUT.with_suffix('.tmp.pptx')
with ZipFile(BASE) as orig,ZipFile(OUT) as edited,ZipFile(temp,'w',ZIP_DEFLATED) as out:
 for f in edited.infolist():
  data=edited.read(f.filename)
  if f.filename.startswith('ppt/slides/slide') and f.filename.endswith('.xml'):
   number=int(f.filename.split('slide')[-1].split('.')[0])
   if number not in changed:data=orig.read(f.filename)
  out.writestr(f,data)
temp.replace(OUT)
print('Updated slides:',sorted(changed));print('PASS: original background objects and existing cell wording preserved; other slide XML unchanged.')
